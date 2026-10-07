"""Excel organizado (mesma estrutura do botão "Baixar Excel organizado" da newsletter).

Abas: Oportunidades (com Responsável/Andamento para preencher), Roteiros (passo a passo de
cada ficha), Contatos (canais oficiais das empresas) e Cadastro (etapas por empresa).
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TEAL, INK, SOFT, LINK = "1F5F5A", "233C38", "65736F", "166A9B"
ZEBRA, EDIT = "F0F5F3", "FFF3CD"
STATUS = ["Não iniciado", "Em contato", "Em cadastro", "Homologado", "Em cotação", "Proposta enviada",
          "Contratado", "Encerrado"]
URL_RE = re.compile(r"^https?://")


def _rows(docs: list[tuple[dict, dict]], accounts: dict):
    """docs = [(linha da notícia, ficha)]"""
    opp, detail, used = [], [], set()
    for row, d in docs:
        nid = str(d.get("newsletter_id") or row.get("id"))
        title = row.get("title") or ""
        accs = [k for k in d.get("accounts") or [] if k in accounts]
        used.update(accs)
        opp.append([nid, title, (row.get("ufs") or "").replace(",", ", "),
                    "\n".join(accounts[k]["name"] for k in accs) or "Identificar contratante",
                    d.get("phase"), d.get("priority"), d.get("need"), (d.get("steps") or [""])[0], "", "Não iniciado"])

        def line(topic, order, text, source=None, timing=None):
            timing = timing or {}
            detail.append([nid, title, topic, order, text, timing.get("internal", ""), timing.get("external", ""),
                           timing.get("start", ""), timing.get("done", ""), source or row.get("url")])

        line("Fase", 1, d.get("phaseDetail"))
        line("Entrada", 1, d.get("entry"))
        line("Comprador", 1, d.get("buyerTarget"))
        timings = d.get("stepTiming") or []
        for i, t in enumerate(d.get("steps") or []):
            tm = timings[i] if i < len(timings) else {}
            line("Roteiro", i + 1, t, (tm.get("reference") or {}).get("url"), tm)
        line("Resultado da primeira ação", 1, d.get("firstOutcome"))
        for i, t in enumerate(d.get("questions") or []):
            line("Perguntas", i + 1, t)
        line("Mensagem", 1, d.get("message"))
        for i, t in enumerate(d.get("gaps") or []):
            line("Pendências", i + 1, t)
        line("Acompanhamento", 1, d.get("followup"))
        line("Critério para avançar", 1, d.get("trigger"))
        for i, c in enumerate(d.get("companies") or []):
            line("Empresas envolvidas", i + 1, f"{c.get('name')} — {c.get('role')} — {c.get('status')}")
        line("Empresas cotadas", 1, d.get("bidStatus"))
        for i, s in enumerate(d.get("sources") or []):
            line("Fonte adicional", i + 1, s.get("name"), s.get("url"))

    contacts, register = [], []
    for k in sorted(used, key=lambda k: accounts[k].get("name", "")):
        a = accounts[k]
        contacts.append([a.get("name"), a.get("about"), a.get("channel"), a.get("kind"),
                         a.get("portal") or "Solicitar procedimento oficial", a.get("url")])
        for i, t in enumerate(a.get("guide") or []):
            register.append([a.get("name"), i + 1, t, a.get("portal") or a.get("url"), a.get("url")])
    return opp, detail, contacts, register


def workbook_bytes(docs: list[tuple[dict, dict]], accounts: dict) -> bytes:
    opp, detail, contacts, register = _rows(docs, accounts)
    sheets = [
        ("Oportunidades", "Radar Cordeiro — oportunidades",
         "Filtre a lista. Use o ID para consultar Roteiros. Responsável e andamento são campos para preencher.",
         ["ID", "Notícia / projeto", "UF", "Conta / canal inicial", "Fase do projeto", "Prioridade",
          "Serviço potencial", "Primeira ação", "Responsável", "Andamento"],
         opp, [14, 49, 12, 29, 30, 17, 40, 60, 22, 24], {8, 9}),
        ("Roteiros", "Roteiros por oportunidade",
         "Filtre por ID e tópico. Etapas em dias úteis. Faixas internas são estimativas de planejamento, não SLAs. "
         "Espera do cliente separada.",
         ["ID", "Notícia / projeto", "Tópico", "Passo", "Orientação", "Tempo interno estimado",
          "Espera / fatores externos", "Quando iniciar", "Critério de avanço", "Fonte de contexto"],
         detail, [14, 49, 23, 9, 75, 33, 60, 58, 58, 28], set()),
        ("Contatos", "Empresas e contatos",
         "Canais públicos. A coluna Natureza esclarece se o contato é geral ou suporte ao cadastro.",
         ["Empresa", "Atuação", "Contato / canal", "Natureza do contato", "Cadastro / portal", "Fonte oficial"],
         contacts, [30, 55, 53, 59, 45, 40], set()),
        ("Cadastro", "Cadastro por empresa",
         "Filtre por empresa e siga as etapas. Cadastro e homologação não garantem cotação.",
         ["Empresa", "Etapa", "Como proceder", "Abrir cadastro / canal", "Fonte oficial"],
         register, [30, 9, 100, 45, 40], set()),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    top = Alignment(vertical="top", wrap_text=True)
    for index, (name, title, note, headers, rows, widths, editable) in enumerate(sheets):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color=TEAL)
        ws["A2"] = note
        ws["A2"].font = Font(name="Arial", size=11, color=SOFT)
        ws.row_dimensions[1].height, ws.row_dimensions[2].height = 28, 36
        for j, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=TEAL)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[4].height = 34
        for i, r in enumerate(rows, 5):
            height = 26
            for j, v in enumerate(r):
                v = "" if v is None else v
                c = ws.cell(row=i, column=j + 1)
                if isinstance(v, str) and URL_RE.match(v):
                    c.value = "Ver fonte" if "Fonte" in headers[j] else "Abrir cadastro"
                    c.hyperlink = v
                    c.font = Font(name="Arial", size=11, color=LINK, underline="single")
                else:
                    c.value = v
                    c.font = Font(name="Arial", size=11, color=INK)
                c.alignment = top
                if j in editable:
                    c.fill = PatternFill("solid", fgColor=EDIT)
                elif i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=ZEBRA)
                lines = sum(max(1, -(-len(t) // max(1, int(widths[j] * 0.95)))) for t in str(c.value).split("\n"))
                height = max(height, lines * 15 + 10)
            ws.row_dimensions[i].height = min(height, 400)
        last = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A4:{last}{max(4, len(rows) + 4)}"
        ws.freeze_panes = "C5" if index < 2 else "B5"
        if index == 0 and rows:
            dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
            dv.add(f"J5:J{len(rows) + 4}")
            ws.add_data_validation(dv)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
